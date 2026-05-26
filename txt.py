import json
import re
import os
from docx import Document
from docx.shared import Pt


def clean_transcription(text):
    for word in []:
        text = re.sub(rf'\b{word}\b', '', text, flags=re.IGNORECASE)
    text = re.sub(r'\{.*?\}', '', text)
    return re.sub(r'[ \t]+', ' ', text).strip()


def read_file(file_name):
    with open(file_name, "r", encoding="utf-8") as f:
        text = f.read()
    try:
        data = json.loads(text)
        return [data] if isinstance(data, dict) else data
    except json.JSONDecodeError:
        return [{"filename": file_name, "transcription": text}]


def parse_segments(text):
    """
    Splits text into (heading, body) pairs.
    Matches ANY content inside [...] including numbers, symbols, spaces.
    """
    pattern = re.compile(r'\[([^\]]+)\]')
    segments = []
    last_end = 0

    for match in pattern.finditer(text):
        before = text[last_end:match.start()].strip()
        if before:
            if segments:
                segments[-1]["body"] = (segments[-1]["body"] + " " + before).strip()
            else:
                segments.append({"heading": None, "body": before})
        segments.append({"heading": match.group(1).strip(), "body": ""})
        last_end = match.end()

    tail = text[last_end:].strip()
    if tail:
        if segments:
            segments[-1]["body"] = (segments[-1]["body"] + " " + tail).strip()
        else:
            segments.append({"heading": None, "body": tail})

    return segments


def add_formatted_text(doc, text):
    for seg in parse_segments(text):
        para = doc.add_paragraph()
        if seg["heading"]:
            run = para.add_run(seg["heading"] + ":  ")
            run.bold = True
            run.font.size = Pt(12)
        if seg["body"]:
            run = para.add_run(seg["body"])
            run.bold = False
            run.font.size = Pt(12)
        doc.add_paragraph("")  # blank line after each passage


def create_doc(input_file):
    data = read_file(input_file)
    doc = Document()

    title = doc.add_heading("Audio Transcriptions", level=0)
    title.runs[0].font.size = Pt(18)

    for item in data:
        audio_name = item.get("filename", "Unknown File") if isinstance(item, dict) else "Unknown File"
        transcription = item.get("transcription", "") if isinstance(item, dict) else str(item)

        h1 = doc.add_heading("File Name", level=1)
        h1.runs[0].font.size = Pt(14)
        doc.add_paragraph(audio_name)

        h2 = doc.add_heading("Transcription", level=1)
        h2.runs[0].font.size = Pt(14)
        add_formatted_text(doc, clean_transcription(transcription))

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output.docx")
    doc.save(output_path)
    print(f"\nDocument saved: {output_path}")


create_doc(input("Enter input file name: "))
